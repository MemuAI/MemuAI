"""Round 29 phase 2 (r29t-apps): regression tests for bugs found by deep testing of MemuNote,
Quick Access explorer/tools/panels and Sync/safe mode. Each test names the bug it guards."""
import io
import os
import sys
import tempfile
import unittest
import zipfile

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, ROOT)

RU = "Съешь же ещё этих мягких французских булок, да выпей чаю. Привет, мир!\n" * 20


class TextEncodingTests(unittest.TestCase):
    def test_utf16_without_bom_cyrillic_is_detected(self):
        # Bug: Cyrillic UTF-16 without BOM has almost no zero bytes (0x04xx) -> was "binary UTF-8".
        from app.memunote.formats import textio
        for enc in ("utf-16-le", "utf-16-be"):
            info = textio.detect_bytes(RU.encode(enc))
            self.assertEqual(info.encoding, enc)
            self.assertFalse(info.binary)

    def test_no_false_utf16_for_single_byte_codepages(self):
        from app.memunote.formats import textio
        self.assertEqual(textio.detect_bytes(RU.encode("utf-8")).encoding, "utf-8")
        self.assertEqual(textio.detect_bytes(RU.encode("cp1251")).encoding, "cp1251")
        self.assertEqual(textio.detect_bytes(RU.encode("koi8-r")).encoding, "koi8-r")
        self.assertEqual(textio.detect_bytes(b"hello world " * 100).encoding, "utf-8")

    def test_lossy_encode_counts_and_is_fast(self):
        # Bug: a per-character fallback loop made saving a big buffer with one emoji take minutes.
        import time
        from app.memunote.formats import textio
        data, lost = textio.encode("Привет? \U0001F600 мир \U0001F600\U0001F600\n", "cp1251", "crlf")
        self.assertEqual(lost, 3)
        self.assertTrue(data.endswith(b"\r\n"))
        self.assertIn("Привет?".encode("cp1251"), data)
        big = "абв" * 3_000_000 + "\U0001F600"
        start = time.time()
        _data, lost = textio.encode(big, "cp1251")
        self.assertEqual(lost, 1)
        self.assertLess(time.time() - start, 5)


class DocxRoundTripTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        try:
            import docx  # noqa: F401
        except ImportError:
            raise unittest.SkipTest("python-docx not installed")

    def _sample(self):
        import docx
        doc = docx.Document()
        doc.core_properties.author = "Иван Петров"
        section = doc.sections[0]
        section.different_first_page_header_footer = True
        section.header.paragraphs[0].text = "Шапка документа"
        section.footer.paragraphs[0].text = "Подвал"
        section.first_page_header.paragraphs[0].text = "Первая страница"
        doc.add_heading("Заголовок", 1)
        doc.add_paragraph("Текст")
        doc.add_paragraph("один", style="List Bullet")
        doc.add_paragraph("два", style="List Bullet")
        doc.add_paragraph("раз", style="List Number")
        buf = io.BytesIO()
        doc.save(buf)
        return buf.getvalue()

    def test_style_based_lists_are_read_as_lists(self):
        # Bug: "List Bullet"/"List Number" keep numPr in the style, not the paragraph -> shown as plain text.
        from app.memunote.formats.docx import read_docx
        blocks, _meta = read_docx(self._sample())
        kinds = [(b.kind, b.list) for b in blocks if hasattr(b, "kind")]
        self.assertIn(("li", "ul"), kinds)
        self.assertIn(("li", "ol"), kinds)

    def test_save_keeps_headers_footers_and_author(self):
        # Bug: saving a .docx silently dropped headers/footers and replaced the author with "MemuNote".
        import docx
        from app.memunote.formats import richhtml
        from app.memunote.formats.docx import read_docx, write_docx
        original = self._sample()
        blocks, _ = read_docx(original)
        saved = write_docx(richhtml.parse_html(richhtml.to_html(blocks)), "t", original)
        again = docx.Document(io.BytesIO(saved))
        section = again.sections[0]
        self.assertEqual(section.header.paragraphs[0].text, "Шапка документа")
        self.assertEqual(section.footer.paragraphs[0].text, "Подвал")
        self.assertTrue(section.different_first_page_header_footer)
        self.assertEqual(section.first_page_header.paragraphs[0].text, "Первая страница")
        self.assertEqual(again.core_properties.author, "Иван Петров")
        texts = [p.text for p in again.paragraphs]
        self.assertIn("один", texts)          # no trailing space from "<li>один\n</li>"
        with zipfile.ZipFile(io.BytesIO(saved)) as archive:
            self.assertIsNone(archive.testzip())

    def test_new_document_without_original_still_valid(self):
        import docx
        from app.memunote.formats import richhtml
        from app.memunote.formats.docx import write_docx
        saved = write_docx(richhtml.parse_html("<p>a</p><ul><li>b</li></ul>"), "t", None)
        self.assertEqual([p.text for p in docx.Document(io.BytesIO(saved)).paragraphs], ["a", "b"])


class RichFormatTests(unittest.TestCase):
    BLOCKS_HTML = "<ul><li>a</li><li>b</li></ul><ol><li>one</li><li>two</li></ol>"

    def test_odt_bullets_then_numbers_stay_two_lists(self):
        # Bug: adjacent <ul> and <ol> were written as ONE text:list with the first list's style -> "1. 2." became bullets.
        from app.memunote.formats import richhtml
        from app.memunote.formats.odt import read_odt, write_odt
        data = write_odt(richhtml.parse_html(self.BLOCKS_HTML), "t")
        blocks, _ = read_odt(data)
        self.assertEqual([b.list for b in blocks if getattr(b, "kind", "") == "li"], ["ul", "ul", "ol", "ol"])

    def test_rtf_bullets_then_numbers_stay_two_lists(self):
        from app.memunote.formats import richhtml
        from app.memunote.formats.rtf import read_rtf, write_rtf
        data = write_rtf(richhtml.parse_html(self.BLOCKS_HTML))
        self.assertEqual(data.count(b"\\listoverride\\listid"), 2)
        blocks, _ = read_rtf(data)
        self.assertEqual([b.list for b in blocks if getattr(b, "kind", "") == "li"], ["ul", "ul", "ol", "ol"])

    def test_rtf_picture_with_picprop_and_header_picture(self):
        # Bug 1: text of {\\*\\picprop …} inside \\pict was read as picture hex -> corrupt PNG.
        # Bug 2: a logo in {\\header …} appeared in the first body paragraph.
        import base64
        from app.memunote.formats.rtf import read_rtf
        png = base64.b64decode("iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVR42mP8z8BQDwAEhQGAhKmMIQAAAABJRU5ErkJggg==")
        pict = "{\\pict{\\*\\picprop{\\sp{\\sn wzName}{\\sv Picture 1}}}\\pngblip\\picwgoal300\\pichgoal300 " + png.hex() + "}"
        rtf = ("{\\rtf1\\ansi{\\header{\\pard " + pict + "\\par}}{\\pard Body " + pict + "\\par}}").encode("ascii")
        blocks, _ = read_rtf(rtf)
        images = [r.image for b in blocks for r in getattr(b, "runs", []) if r.image]
        self.assertEqual(len(images), 1)
        self.assertEqual(base64.b64decode(images[0].split(",", 1)[1]), png)


if __name__ == "__main__":
    unittest.main()
