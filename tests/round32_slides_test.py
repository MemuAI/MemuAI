"""Round 32, zone "slides": presentations in MemuNote (app/memunote/static/js/slides/*, app/memunote/slides_api.py).

Static / unit tests always run. Browser tests (real MemuNote server + headless Chromium) skip without
Playwright or Chromium; python-pptx (validation oracle) and pypdf checks skip when the package is missing.
The live AI test needs XKIRO_BASE / XKIRO_KEY (source /workspace/work/secrets/test_keys.env), otherwise skips.

Run: NO_PROXY=127.0.0.1,localhost /workspace/work/venv/bin/python -m pytest tests/round32_slides_test.py -q
"""
from __future__ import annotations

import http.client
import http.server
import io
import json
import os
import re
import shutil
import subprocess
import sys
import tempfile
import threading
import time
import unittest
import urllib.request
import zipfile
from pathlib import Path
from unittest import mock

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
STATIC = ROOT / "app" / "memunote" / "static"
SLIDES = STATIC / "js" / "slides"
os.environ.setdefault("NO_PROXY", "127.0.0.1,localhost")

try:
    import pptx  # noqa: F401 - validation oracle
    HAVE_PPTX = True
except ImportError:
    HAVE_PPTX = False


def _chromium():
    for path in ("/usr/bin/chromium", "/workspace/work/chromium/chrome", shutil.which("chromium") or "", shutil.which("google-chrome") or ""):
        if path and os.path.exists(path):
            return path
    return None


def package_problems(data: bytes):
    """Structural OOXML checks: XML well-formed, every part typed, no dangling rels, unique shape ids per slide."""
    import posixpath
    import xml.etree.ElementTree as ET
    z = zipfile.ZipFile(io.BytesIO(data))
    names = set(z.namelist())
    probs = []
    ct = ET.fromstring(z.read("[Content_Types].xml"))
    ns = "{http://schemas.openxmlformats.org/package/2006/content-types}"
    over = {o.get("PartName").lstrip("/") for o in ct.iter(ns + "Override")}
    defs = {d.get("Extension").lower() for d in ct.iter(ns + "Default")}
    for n in names:
        if n.endswith(".xml") or n.endswith(".rels"):
            try:
                ET.fromstring(z.read(n))
            except ET.ParseError as exc:
                probs.append(f"bad xml {n}: {exc}")
        if n not in over and n.rsplit(".", 1)[-1].lower() not in defs:
            probs.append(f"no content type: {n}")
    probs += [f"override without part: {o}" for o in over if o not in names]
    for n in names:
        if not n.endswith(".rels"):
            continue
        base = posixpath.dirname(posixpath.dirname(n))
        for r in ET.fromstring(z.read(n)):
            if r.get("TargetMode") == "External":
                continue
            t = r.get("Target")
            p = t.lstrip("/") if t.startswith("/") else posixpath.normpath(posixpath.join(base, t))
            if p not in names:
                probs.append(f"{n}: dangling {t}")
    for n in names:
        if re.match(r"ppt/slides/slide\d+\.xml$", n):
            ids = re.findall(rb'<p:cNvPr id="(\d+)"', z.read(n))
            if len(ids) != len(set(ids)):
                probs.append(f"duplicate shape ids in {n}")
    return probs


# ============================================================================ static / unit
class StaticTest(unittest.TestCase):
    def test_i18n_keys_match_and_english_has_no_cyrillic(self) -> None:
        text = (SLIDES / "i18n-slides.js").read_text(encoding="utf-8")
        ru = text[text.index("  ru: {"):text.index("  en: {")]
        en = text[text.index("  en: {"):]
        keys = lambda block: set(re.findall(r"\b([a-z][a-z0-9_]*): \"", block))
        self.assertEqual(keys(ru), keys(en))
        self.assertFalse(re.search(r"[А-Яа-яЁё]", en), "Cyrillic in the English dictionary")
        used = set()
        for f in SLIDES.glob("*.js"):
            used |= {k for k in re.findall(r't\("((?:sl|sh)_[a-z0-9_]+)"', f.read_text(encoding="utf-8")) if not k.endswith("_")}
        self.assertFalse(used - keys(en), "missing i18n keys")

    def test_wiring_in_shared_files(self) -> None:
        app_js = (STATIC / "js" / "app.js").read_text(encoding="utf-8")
        self.assertIn('from "./slides/view.js"', app_js)
        self.assertIn('tab.kind === "slides"', app_js)
        self.assertIn("slidesOwnsKey(e)", app_js)
        self.assertIn("slides.css", (STATIC / "index.html").read_text(encoding="utf-8"))
        self.assertIn("SLIDES_I18N", (STATIC / "js" / "i18n.js").read_text(encoding="utf-8"))
        self.assertIn("app.memunote.slides_api", (ROOT / "build.py").read_text(encoding="utf-8"))
        self.assertIn("app.memunote.slides_api", (ROOT / "MemuNote.spec").read_text(encoding="utf-8"))

    def test_spec_imports_check(self) -> None:
        r = subprocess.run([sys.executable, "tests/round29_spec_imports_check.py"], cwd=ROOT, capture_output=True, text=True, timeout=300)
        mine = [ln for ln in (r.stdout + r.stderr).splitlines() if ln.startswith("PROBLEM") and ("memunote" in ln.lower() or "slides" in ln)]
        self.assertEqual(mine, [], r.stdout[-2000:])      # problems of other zones are not this zone's business

    def test_kinds_and_associations(self) -> None:
        from app.memunote import server, winsys
        for ext in (".pptx", ".PPTX", ".odp", ".potx", ".ppsx", ".pptm"):
            self.assertEqual(server.kind_of("C:/a/deck" + ext), "slides")
        self.assertEqual(server.kind_of("C:/a/old.ppt"), "legacy")
        self.assertIn(".pptx", winsys.SAFE_EXT)
        self.assertIn(".odp", winsys.SAFE_EXT)
        self.assertNotIn(".pptx", winsys.NEVER_EXT)

    def test_render_css_block_exists(self) -> None:
        css = (STATIC / "slides.css").read_text(encoding="utf-8")
        self.assertLess(css.index("/* @render-start */"), css.index("/* @render-end */"))

    def test_js_syntax(self) -> None:
        node = shutil.which("node")
        if not node:
            self.skipTest("node is not installed")
        for f in sorted(SLIDES.glob("*.js")):
            r = subprocess.run([node, "--check", str(f)], capture_output=True, text=True)
            self.assertEqual(r.returncode, 0, f"{f.name}: {r.stderr}")


class PromptTest(unittest.TestCase):
    def test_build_messages(self) -> None:
        from app.memunote import slides_api
        m = slides_api.build_messages("generate", {"topic": "Кошки <script>", "count": 7, "lang": "ru", "source": "SRC"}, "ru")
        self.assertEqual([x["role"] for x in m], ["system", "user"])
        self.assertIn("Russian", m[0]["content"])
        self.assertIn("exactly 7 slides", m[0]["content"])
        self.assertNotIn("Кошки", m[0]["content"])          # user data never in the system prompt
        self.assertIn("<source>\nSRC\n</source>", m[1]["content"])
        m = slides_api.build_messages("translate", {"target": "de", "items": [{"id": "1", "text": "a"}]}, "ru")
        self.assertIn("German", m[0]["content"])
        for task in ("rewrite", "layout", "notes"):
            self.assertIn("JSON", slides_api.build_messages(task, {"items": [], "slides": []}, "en")[0]["content"])
        with self.assertRaises(ValueError):
            slides_api.build_messages("rm -rf", {}, "en")


class _ServerCase(unittest.TestCase):
    def setUp(self) -> None:
        self.tmp = tempfile.TemporaryDirectory(prefix="r32sl-")
        self.patch = mock.patch.dict(os.environ, {"MEMUAI_DATA_DIR": os.path.join(self.tmp.name, "data")})
        self.patch.start()
        from app.memunote import server
        self.server = server
        self.app = server.NoteApp([])
        self.httpd, _ = server.start(self.app)
        self.dir = Path(self.tmp.name)

    def tearDown(self) -> None:
        self.httpd.shutdown()
        self.httpd.server_close()
        self.patch.stop()
        self.tmp.cleanup()

    def call(self, method, path, body=None, raw=False, headers=None):
        conn = http.client.HTTPConnection("127.0.0.1", self.app.port, timeout=60)
        hdrs = {"X-MN-Token": self.app.token, "Host": f"127.0.0.1:{self.app.port}"}
        if isinstance(body, (dict, list)):
            body = json.dumps(body).encode()
            hdrs["Content-Type"] = "application/json"
        hdrs.update(headers or {})
        conn.request(method, path, body=body, headers=hdrs)
        resp = conn.getresponse()
        data = resp.read()
        conn.close()
        return resp, (data if raw else json.loads(data or b"null"))


class ApiTest(_ServerCase):
    def test_open_and_save_binary_guard(self) -> None:
        target = self.dir / "deck.pptx"
        buf = io.BytesIO()
        with zipfile.ZipFile(buf, "w") as z:
            z.writestr("[Content_Types].xml", "<Types/>")
        target.write_bytes(buf.getvalue())
        _, info = self.call("POST", "/api/open", {"path": str(target)})
        self.assertEqual(info["kind"], "slides")
        self.assertNotIn("text", info)
        import urllib.parse
        q = urllib.parse.quote(str(target))
        resp, data = self.call("POST", f"/api/save_binary?path={q}&force=1", b"not a zip")
        self.assertEqual(resp.status, 400)
        self.assertEqual(target.read_bytes(), buf.getvalue())   # untouched
        resp, data = self.call("POST", f"/api/save_binary?path={q}&force=1", buf.getvalue() + b"")
        self.assertTrue(data["ok"])

    def test_slides_ai_route_with_fake_proxy(self) -> None:
        answer = json.dumps({"title": "T", "slides": [{"layout": "title", "title": "Привет", "subtitle": "мир"}]}, ensure_ascii=False)
        fake = _start_fake_proxy({"generate": answer})
        with mock.patch.dict(os.environ, {"MEMUAI_PROXY_URL": f"http://127.0.0.1:{fake.server_address[1]}"}):
            _, raw = self.call("POST", "/api/slides_ai", {"task": "generate", "model": "fake", "input": {"topic": "x", "count": 3}}, raw=True)
        fake.shutdown()
        events = [json.loads(line[5:]) for line in raw.decode("utf-8").split("\n") if line.startswith("data:")]
        self.assertEqual(events[0]["t"], "start")
        self.assertEqual(events[-1]["t"], "done")
        self.assertEqual(json.loads("".join(e["text"] for e in events if e["t"] == "delta")), json.loads(answer))
        resp, data = self.call("POST", "/api/slides_ai", {"task": "nope", "model": "fake", "input": {}})
        self.assertEqual(resp.status, 400)


# ============================================================================ fake proxy for AI flows
FAKE_DECK = {"title": "Кофе", "slides": [
    {"layout": "title", "title": "Кофе: история и культура", "subtitle": "Короткий обзор", "notes": "Приветствие."},
    {"layout": "content", "title": "Откуда кофе", "bullets": ["Эфиопия, IX век", "  Легенда о пастухе Калди", "Йемен и суфии"], "notes": "Рассказать легенду."},
    {"layout": "two", "title": "Арабика и робуста", "left": ["Мягкий вкус", "Больше кислотности"], "right": ["Крепче", "Больше кофеина"]},
    {"layout": "chart", "title": "Производство", "chart": {"type": "col", "cats": ["Бразилия", "Вьетнам"], "series": [{"name": "млн мешков", "vals": [66, 29]}]}},
    {"layout": "table", "title": "Сорта", "table": [["Сорт", "Страна"], ["Сидамо", "Эфиопия"]]},
    {"layout": "section", "title": "Итоги", "subtitle": "Вопросы?"},
]}


class _FakeProxy(http.server.BaseHTTPRequestHandler):
    answers: dict = {}

    def log_message(self, *a):  # noqa: D401
        pass

    def do_GET(self):  # noqa: N802
        body = json.dumps({"service": "memuai-proxy"} if self.path == "/health" else {"models": [{"id": "fake", "label": "Fake"}], "default_model": "fake"}).encode()
        self.send_response(200)
        self.send_header("Content-Type", "application/json")
        self.send_header("Content-Length", str(len(body)))
        self.end_headers()
        self.wfile.write(body)

    def do_POST(self):  # noqa: N802
        req = json.loads(self.rfile.read(int(self.headers.get("Content-Length") or 0)) or b"{}")
        system = req.get("messages", [{}])[0].get("content", "")
        user = req.get("messages", [{}, {}])[-1].get("content", "")
        if "presentation designer and writer" in system:
            ans = self.answers.get("generate", json.dumps(FAKE_DECK, ensure_ascii=False))
        elif "Translate every text" in system:
            items = json.loads(user[user.index("{"):user.rindex("}") + 1])["items"]
            ans = "```json\n" + json.dumps({"items": [{"id": i["id"], "text": "EN " + i["text"]} for i in items]}, ensure_ascii=False) + "\n```"
        elif "speaker notes" in system.lower():
            slides = json.loads(user[user.index("["):user.rindex("]") + 1])
            ans = json.dumps({"notes": [{"n": s["n"], "text": f"Заметка {s['n']}"} for s in slides]}, ensure_ascii=False)
        elif "text boxes of one presentation slide" in system:
            items = json.loads(user[user.index("{"):user.rindex("}") + 1])["items"]
            ans = "Вот результат: " + json.dumps({"items": [{"id": i["id"], "text": i["text"].split("\n")[0][:10]} for i in items]}, ensure_ascii=False)
        else:
            ans = self.answers.get("default", "{}")
        self.send_response(200)
        self.send_header("Content-Type", "text/event-stream")
        self.end_headers()
        for i in range(0, len(ans), 7):
            self.wfile.write(("data: " + json.dumps({"choices": [{"delta": {"content": ans[i:i + 7]}}]}, ensure_ascii=False) + "\n\n").encode())
        self.wfile.write(b"data: [DONE]\n\n")


def _start_fake_proxy(answers=None):
    _FakeProxy.answers = answers or {}
    fake = http.server.ThreadingHTTPServer(("127.0.0.1", 0), _FakeProxy)
    threading.Thread(target=fake.serve_forever, daemon=True).start()
    return fake


def make_sample_pptx(path: Path, slides: int = 3) -> None:
    """A third-party deck (python-pptx's default template): chart with embedded workbook, table, group, notes."""
    from pptx import Presentation
    from pptx.chart.data import CategoryChartData
    from pptx.dml.color import RGBColor
    from pptx.enum.chart import XL_CHART_TYPE
    from pptx.util import Inches
    p = Presentation()
    s = p.slides.add_slide(p.slide_layouts[0]); s.shapes.title.text = "Отчёт"; s.placeholders[1].text = "Подзаголовок"
    s = p.slides.add_slide(p.slide_layouts[1]); s.shapes.title.text = "Итоги"
    tf = s.placeholders[1].text_frame; tf.text = "Рост 12%"; para = tf.add_paragraph(); para.text = "Клиенты"; para.level = 1
    r = tf.add_paragraph().add_run(); r.text = "Красный"; r.font.bold = True; r.font.color.rgb = RGBColor(0xC0, 0, 0)
    s.notes_slide.notes_text_frame.text = "Заметка 2"
    s = p.slides.add_slide(p.slide_layouts[5]); s.shapes.title.text = "Данные"
    cd = CategoryChartData(); cd.categories = ["Q1", "Q2"]; cd.add_series("A", (1, 3))
    s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(0.5), Inches(1.5), Inches(4), Inches(3), cd)
    tb = s.shapes.add_table(2, 2, Inches(5), Inches(1.5), Inches(4), Inches(1)).table
    tb.cell(0, 0).text = "H"; tb.cell(1, 1).text = "V"
    g = s.shapes.add_group_shape(); g.shapes.add_shape(9, Inches(1), Inches(5), Inches(1), Inches(1)); g.shapes.add_shape(1, Inches(2.5), Inches(5), Inches(1), Inches(1))
    for i in range(3, slides):
        s = p.slides.add_slide(p.slide_layouts[1]); s.shapes.title.text = f"Слайд {i + 1}"; s.placeholders[1].text = f"Пункт {i + 1}"
    p.save(str(path))


def make_sample_odp(path: Path) -> None:
    content = """<?xml version="1.0" encoding="UTF-8"?>
<office:document-content xmlns:office="urn:oasis:names:tc:opendocument:xmlns:office:1.0" xmlns:style="urn:oasis:names:tc:opendocument:xmlns:style:1.0"
 xmlns:text="urn:oasis:names:tc:opendocument:xmlns:text:1.0" xmlns:draw="urn:oasis:names:tc:opendocument:xmlns:drawing:1.0"
 xmlns:fo="urn:oasis:names:tc:opendocument:xmlns:xsl-fo-compatible:1.0" xmlns:svg="urn:oasis:names:tc:opendocument:xmlns:svg-compatible:1.0"
 xmlns:presentation="urn:oasis:names:tc:opendocument:xmlns:presentation:1.0" office:version="1.2">
<office:automatic-styles><style:style style:name="T1" style:family="text"><style:text-properties fo:font-weight="bold" fo:color="#c00000"/></style:style>
<style:style style:name="gr1" style:family="graphic"><style:graphic-properties draw:fill="solid" draw:fill-color="#2f6feb"/></style:style></office:automatic-styles>
<office:body><office:presentation>
<draw:page draw:name="p1"><draw:frame presentation:class="title" svg:x="2cm" svg:y="1cm" svg:width="24cm" svg:height="3cm"><draw:text-box><text:p>ODP заголовок</text:p></draw:text-box></draw:frame>
<draw:frame svg:x="2cm" svg:y="5cm" svg:width="20cm" svg:height="8cm"><draw:text-box><text:list><text:list-item><text:p>Пункт <text:span text:style-name="T1">жирный</text:span></text:p></text:list-item></text:list></draw:text-box></draw:frame>
<draw:custom-shape draw:style-name="gr1" svg:x="20cm" svg:y="12cm" svg:width="4cm" svg:height="3cm"><text:p>Фигура</text:p><draw:enhanced-geometry draw:type="ellipse"/></draw:custom-shape>
<presentation:notes><draw:frame><draw:text-box><text:p>Заметка ODP</text:p></draw:text-box></draw:frame></presentation:notes></draw:page>
<draw:page draw:name="p2"><draw:frame svg:x="2cm" svg:y="2cm" svg:width="10cm" svg:height="3cm"><draw:text-box><text:p>Второй</text:p></draw:text-box></draw:frame></draw:page>
</office:presentation></office:body></office:document-content>"""
    styles = """<?xml version="1.0" encoding="UTF-8"?><office:document-styles xmlns:office="urn:oasis:names:tc:opendocument:xmlns:office:1.0" xmlns:style="urn:oasis:names:tc:opendocument:xmlns:style:1.0"
 xmlns:fo="urn:oasis:names:tc:opendocument:xmlns:xsl-fo-compatible:1.0"><office:automatic-styles><style:page-layout style:name="PM1"><style:page-layout-properties fo:page-width="28cm" fo:page-height="15.75cm"/></style:page-layout></office:automatic-styles></office:document-styles>"""
    with zipfile.ZipFile(path, "w") as z:
        z.writestr(zipfile.ZipInfo("mimetype"), "application/vnd.oasis.opendocument.presentation")
        z.writestr("content.xml", content)
        z.writestr("styles.xml", styles)


# ============================================================================ browser
@unittest.skipUnless(_chromium(), "chromium is not available")
class BrowserTest(_ServerCase):
    """Real MemuNote page: the deck model / OOXML writer in the page, and the editor UI."""

    def setUp(self) -> None:
        try:
            from playwright.sync_api import sync_playwright
        except ImportError:
            self.skipTest("playwright is not installed")
        self.fake = _start_fake_proxy()
        self.env = mock.patch.dict(os.environ, {"MEMUAI_PROXY_URL": f"http://127.0.0.1:{self.fake.server_address[1]}"})
        self.env.start()
        super().setUp()
        self.call("POST", "/api/settings", {"language": "en", "restore_session": False, "ai_model": "fake"})
        self.pw = sync_playwright().start()
        self.browser = self.pw.chromium.launch(executable_path=_chromium(), args=["--no-sandbox"])
        self.pg = self.browser.new_context(viewport={"width": 1280, "height": 800}, permissions=["clipboard-read", "clipboard-write"]).new_page()   # Round 34: Ctrl+C/Ctrl+V need the clipboard in headless
        self.errors = []
        self.pg.on("pageerror", lambda e: self.errors.append(str(e)))
        self.pg.goto(f"http://127.0.0.1:{self.app.port}/index.html#t={self.app.token}")
        self.pg.wait_for_selector("body.ready", timeout=20000)

    def tearDown(self) -> None:
        try:
            self.browser.close()
            self.pw.stop()
        finally:
            super().tearDown()
            self.env.stop()
            self.fake.shutdown()

    def open(self, path, wait=".sl-thumb"):
        self.call("POST", "/api/remote", {"action": "open", "paths": [str(path)]})
        self.pg.wait_for_selector(wait, timeout=20000)
        time.sleep(0.6)

    def view(self, expr):
        return self.pg.evaluate("() => { const v = MN.tabs.find(t => t.kind === 'slides' && t.box && !t.box.hidden)?.comp; return (" + expr + ") }")

    def js(self, body):
        return self.pg.evaluate("async () => { " + body + " }")

    def test_model_roundtrip_and_validation(self) -> None:
        out = self.dir / "model.pptx"
        r = self.js("""
          const P = await import('/js/slides/pptx.js');
          const deck = await P.newDeck({ lang: 'ru' });
          deck.slides[0].elements[0].text = { paras: [{ runs: [{ t: 'Заголовок' }] }] };
          const s = P.makeSlide(deck, deck.layouts[1].path);
          s.elements[0].text = { paras: [{ runs: [{ t: 'Second', b: true, color: '#FF0000', sz: 40 }] }] };
          s.elements[1].text = { paras: [{ runs: [{ t: 'one' }] }, { lvl: 1, runs: [{ t: 'two' }] }] };
          s.elements.push({ id: 'o', kind: 'shape', geom: 'ellipse', x: 1e6, y: 1e6, w: 2e6, h: 1e6, rot: 30, fill: { color: 'scheme:accent2' }, line: { w: 25400, color: '#000000' }, text: { paras: [{ runs: [{ t: 'Oval' }] }] }, name: 'Oval' });
          s.elements.push({ id: 't', kind: 'table', x: 1e6, y: 3e6, w: 6e6, h: 8e5, name: 'Table', table: { cols: [3e6, 3e6], firstRow: true, bandRow: true, rows: [{ h: 4e5, cells: [{ text: { paras: [{ runs: [{ t: 'A' }] }] } }, { text: { paras: [{ runs: [{ t: 'B' }] }] } }] }, { h: 4e5, cells: [{ text: { paras: [{ runs: [{ t: '1' }] }] } }, { text: { paras: [{ runs: [{ t: '2' }] }] } }] }] } });
          s.elements.push({ id: 'c', kind: 'chart', x: 7e6, y: 3e6, w: 4e6, h: 3e6, name: 'Chart', chart: { type: 'pie', cats: ['x', 'y'], series: [{ name: 'S', vals: [3, 5] }] } });
          const png = Uint8Array.from(atob('iVBORw0KGgoAAAANSUhEUgAAAAIAAAACCAYAAABytg0kAAAAFklEQVQIW2P8z8Dwn4GBgYGJAQoAAP//GQQBAX6yKe8AAAAASUVORK5CYII='), c => c.charCodeAt(0));
          s.elements.push({ id: 'p', kind: 'pic', x: 9e6, y: 5e5, w: 1e6, h: 1e6, name: 'Picture', pic: { media: deck.addMedia(png, 'image/png'), crop: { l: 10, t: 0, r: 0, b: 5 } } });
          s.notes = 'Заметки докладчика'; s.transition = { type: 'push', dir: 'u', spd: 'fast' };
          s.elements[2].anim = { type: 'fly', dir: 'l', trig: 'click', dur: 600, order: 0 };
          s.elements[3].anim = { type: 'fade', trig: 'with', dur: 400, order: 1 };
          s.hidden = true;
          deck.slides.push(s);
          const bytes = await P.savePptx(deck);
          const d2 = await P.openPptx(bytes);
          const s2 = d2.slides[1];
          let bin = ''; for (const b of bytes) bin += String.fromCharCode(b);
          return { b64: btoa(bin), n: d2.slides.length, title: P.plainText(d2.slides[0].elements[0].text), body: s2.elements[1].text.paras.map(p => (p.lvl || 0) + p.runs.map(r => r.t).join('')),
            run: s2.elements[0].text.paras[0].runs[0], notes: s2.notes, tr: s2.transition, anims: s2.elements.map(e => e.anim ? e.anim.type + ':' + e.anim.trig : null), hidden: s2.hidden,
            chart: s2.elements[4].chart, crop: s2.elements[5].pic.crop, rot: s2.elements[2].rot, table: s2.elements[3].table.rows.map(r => r.cells.map(c => P.plainText(c.text))) };
        """)
        import base64
        data = base64.b64decode(r.pop("b64"))
        out.write_bytes(data)
        self.assertEqual(package_problems(data), [])
        self.assertEqual(r["n"], 2)
        self.assertEqual(r["title"], "Заголовок")
        self.assertEqual(r["body"], ["0one", "1two"])
        self.assertTrue(r["run"]["b"]); self.assertEqual(r["run"]["sz"], 40); self.assertEqual(r["run"]["color"], "#FF0000")
        self.assertEqual(r["notes"], "Заметки докладчика")
        self.assertEqual((r["tr"]["type"], r["tr"]["dir"]), ("push", "u"))
        self.assertEqual(r["anims"][2:4], ["fly:click", "fade:with"])
        self.assertTrue(r["hidden"])
        self.assertEqual((r["chart"]["type"], r["chart"]["series"][0]["vals"]), ("pie", [3, 5]))
        self.assertEqual(r["crop"], {"l": 10, "t": 0, "r": 0, "b": 5})
        self.assertEqual(r["rot"], 30)
        self.assertEqual(r["table"], [["A", "B"], ["1", "2"]])
        if HAVE_PPTX:
            from pptx import Presentation
            prs = Presentation(str(out))
            self.assertEqual(len(prs.slides), 2)
            texts = [sh.text_frame.text for sh in prs.slides[1].shapes if sh.has_text_frame]
            self.assertIn("Second", texts)
            self.assertEqual(prs.slides[1].notes_slide.notes_text_frame.text, "Заметки докладчика")
            kinds = {sh.shape_type for sh in prs.slides[1].shapes}
            self.assertTrue(any(sh.has_chart for sh in prs.slides[1].shapes if hasattr(sh, "has_chart")))
            self.assertTrue(any(sh.has_table for sh in prs.slides[1].shapes if hasattr(sh, "has_table")))
            self.assertTrue(kinds)

    @unittest.skipUnless(HAVE_PPTX, "python-pptx is not installed")
    def test_third_party_deck_is_preserved_byte_for_byte(self) -> None:
        src = self.dir / "third.pptx"
        make_sample_pptx(src)
        orig = self.dir / "orig.pptx"
        shutil.copy(src, orig)
        self.open(src)
        # edit one run on slide 2 and save through the UI (Ctrl+S → /api/save_binary)
        self.view("(() => { const s = v.deck.slides[1]; s.elements[0].text.paras[0].runs[0].t = 'Итоги (новое)'; v.touch(s); v.snapshot(); return 1 })()")
        self.pg.keyboard.press("Control+s")
        time.sleep(1.5)
        new = src.read_bytes()
        self.assertEqual(package_problems(new), [])
        z_new = zipfile.ZipFile(io.BytesIO(new))
        z_old = zipfile.ZipFile(orig)
        # every part except the edited slide is identical ([Content_Types].xml / presentation.xml too: re-serialized only when really changed)
        changed = [n for n in z_old.namelist() if n in z_new.namelist() and z_old.read(n) != z_new.read(n)]
        self.assertEqual(sorted(set(changed) - {"ppt/slides/slide2.xml", "docProps/core.xml"}), [])
        self.assertEqual(sorted(set(z_old.namelist()) - set(z_new.namelist())), [])
        from pptx import Presentation
        prs = Presentation(str(src))
        self.assertEqual(prs.slides[1].shapes.title.text, "Итоги (новое)")
        body = prs.slides[1].placeholders[1].text_frame
        self.assertEqual([p.text for p in body.paragraphs], ["Рост 12%", "Клиенты", "Красный"])
        self.assertTrue(body.paragraphs[2].runs[0].font.bold)
        self.assertEqual(prs.slides[1].notes_slide.notes_text_frame.text, "Заметка 2")
        self.assertTrue(prs.slides[2].shapes[1].has_chart)
        self.assertEqual(self.errors, [])

    def test_editor_flow_new_deck_type_move_undo_save(self) -> None:
        out = self.dir / "flow.pptx"
        pg = self.pg
        pg.click("#btn-menu"); pg.click(".menu .mi:has-text('New presentation')")
        pg.click(".sl-new-th[data-th='ocean']"); pg.click(".modal-foot .btn.primary")
        pg.wait_for_selector(".sl-thumb")
        els = pg.locator(".sl-canvas .sl-slide > .sl-el:not(.sl-decor)")
        els.nth(0).dblclick(); pg.keyboard.type("My talk"); pg.keyboard.press("Escape")
        pg.keyboard.press("Control+m"); time.sleep(0.3)
        self.assertEqual(self.view("v.deck.slides.length"), 2)
        els.nth(1).dblclick(); pg.keyboard.type("Alpha"); pg.keyboard.press("Enter"); pg.keyboard.type("Beta"); pg.keyboard.press("Tab")
        pg.keyboard.press("Control+b"); pg.keyboard.press("Escape"); pg.keyboard.press("Escape")
        paras = self.view("v.slide.elements[1].text.paras.map(p => [(p.lvl||0), p.runs.map(r => r.t).join(''), !!p.runs.find(r => r.b)])")
        self.assertEqual(paras, [[0, "Alpha", False], [1, "Beta", True]])
        # insert an ellipse, drag it, undo / redo
        pg.click("[data-c='rtab:insert']"); pg.click("[data-c='ins:shape']"); pg.click(".sl-shape-btn[data-g='ellipse']")
        box = pg.locator(".sl-canvas-wrap").bounding_box()
        pg.mouse.click(box["x"] + box["width"] * 0.7, box["y"] + box["height"] * 0.6)
        x0 = self.view("v.selected()[0].x")
        n = els.last.bounding_box()
        pg.mouse.move(n["x"] + n["width"] / 2, n["y"] + n["height"] / 2); pg.mouse.down()
        pg.mouse.move(n["x"] + n["width"] / 2 - 100, n["y"] + n["height"] / 2, steps=6); pg.mouse.up()
        x1 = self.view("v.selected()[0].x")
        self.assertLess(x1, x0)
        pg.keyboard.press("Control+z"); self.assertEqual(self.view("v.slide.elements.at(-1).x"), x0)
        pg.keyboard.press("Control+y"); self.assertEqual(self.view("v.slide.elements.at(-1).x"), x1)
        # duplicate slide, reorder by keyboard in the sorter, delete it again
        pg.click(".sl-thumb[data-i='1']"); pg.keyboard.press("Control+d"); time.sleep(0.3)
        self.assertEqual(self.view("v.deck.slides.length"), 3)
        pg.keyboard.press("Delete"); time.sleep(0.2)
        self.assertEqual(self.view("v.deck.slides.length"), 2)
        # save as (browser mode: path prompt)
        pg.keyboard.press("Control+s"); pg.fill("#mn-prompt", str(out)); pg.keyboard.press("Enter")
        for _ in range(40):
            if out.exists():
                break
            time.sleep(0.2)
        time.sleep(0.5)
        self.assertTrue(out.exists())
        self.assertFalse(self.view("v.dirty"))
        self.assertEqual(package_problems(out.read_bytes()), [])
        if HAVE_PPTX:
            from pptx import Presentation
            prs = Presentation(str(out))
            self.assertIn("Alpha\nBeta", [sh.text_frame.text for sh in prs.slides[1].shapes if sh.has_text_frame])
            self.assertEqual(prs.slides[0].shapes.title.text, "My talk")
        self.assertEqual(self.errors, [])

    def test_odp_import_and_presenter_and_pdf_export(self) -> None:
        odp = self.dir / "talk.odp"
        make_sample_odp(odp)
        self.open(odp)
        self.assertEqual(self.view("v.deck.slides.length"), 2)
        texts = self.view("v.deck.slides[0].elements.map(e => e.text ? e.text.paras.map(p => p.runs.map(r => r.t).join('')).join('|') : e.kind)")
        self.assertIn("ODP заголовок", texts); self.assertIn("Пункт жирный", texts); self.assertIn("Фигура", texts)
        self.assertEqual(self.view("v.deck.slides[0].notes"), "Заметка ODP")
        self.assertEqual(self.view("v.deck.slides[0].elements[1].text.paras[0].runs[1].color"), "#C00000")
        # slide show with presenter view
        self.pg.evaluate("localStorage.setItem('mn_sl_presenter', '1')")
        self.pg.click(".sl-stage"); self.pg.keyboard.press("Escape"); self.pg.keyboard.press("F5")
        self.pg.wait_for_selector(".sl-show.presenter", timeout=5000)
        self.assertIn("Заметка ODP", self.pg.inner_text(".sl-pv-notes"))
        self.pg.keyboard.press("ArrowRight"); time.sleep(0.9)
        self.assertIn("2", self.pg.inner_text(".sl-pv-count"))
        self.pg.keyboard.press("Escape"); time.sleep(0.3)
        self.assertIsNone(self.pg.query_selector(".sl-show"))
        # export PDF (browser mode: the save dialog is a path prompt); ODP saves as .pptx
        pdf = self.dir / "talk.pdf"
        self.pg.click("[data-c='rtab:show']"); self.pg.click("[data-c='exp:pdf']")
        self.pg.fill("#mn-prompt", str(pdf)); self.pg.keyboard.press("Enter")
        for _ in range(60):
            if pdf.exists() and pdf.stat().st_size > 1000:
                break
            time.sleep(0.25)
        self.assertTrue(pdf.read_bytes().startswith(b"%PDF-"))
        try:
            import pypdf
            self.assertEqual(len(pypdf.PdfReader(str(pdf)).pages), 2)
        except ImportError:
            pass
        self.assertEqual(self.errors, [])

    def test_ai_generate_notes_translate_with_fake_proxy(self) -> None:
        self.pg.click("#btn-menu"); self.pg.click(".menu .mi:has-text('New presentation')")
        self.pg.click(".modal-foot .btn:has-text('Create with AI')")
        self.pg.wait_for_selector("#ai-topic", timeout=10000)
        self.pg.fill("#ai-topic", "Coffee")
        self.pg.click(".modal-foot .btn.primary")
        self._wait_js("() => MN.tabs.find(t => t.kind === 'slides')?.comp?.deck?.slides.length === 6", timeout=20000)
        kinds = self.view("v.deck.slides.map(s => s.elements.map(e => e.kind).join(','))")
        self.assertIn("chart", kinds[3]); self.assertIn("table", kinds[4])
        self.assertEqual(self.view("v.deck.slides[1].elements[1].text.paras.map(p => p.lvl || 0)"), [0, 1, 0])
        self.assertEqual(self.view("v.deck.slides[1].notes"), "Рассказать легенду.")
        self.pg.click("[data-c='rtab:ai']"); self.pg.click("[data-c='ai:notesall']")
        self._wait_js("() => MN.tabs.find(t => t.kind === 'slides').comp.deck.slides[5].notes === 'Заметка 6'", timeout=20000)
        self.pg.click("[data-c='ai:translate']"); self.pg.click(".modal-foot .btn.primary")
        self._wait_js("() => MN.tabs.find(t => t.kind === 'slides').comp.deck.slides[0].notes.startsWith('EN ')", timeout=20000)
        self.assertTrue(self.view("v.text()").count("EN ") > 5)
        self.pg.keyboard.press("Control+z"); time.sleep(0.3)        # one undo step reverts the translation
        self.assertFalse(self.view("v.deck.slides[0].notes").startswith("EN "))
        self.assertEqual(self.errors, [])

    # ---------------------------------------------------------------- Round 32 wave 2 (finisher) regressions
    def _new_deck(self):
        self.pg.click("#btn-menu"); self.pg.click(".menu .mi:has-text('New presentation')")
        n = self.pg.evaluate("MN.tabs.filter(t => t.kind === 'slides').length")
        self.pg.click(".sl-new-th[data-th='ocean']"); self.pg.click(".modal-foot .btn.primary")
        self._wait_js(f"() => MN.tabs.filter(t => t.kind === 'slides' && t.comp?.deck).length > {n}", timeout=20000); time.sleep(0.4)

    def test_overflowing_bullets_never_overlap_and_master_autofit_is_inherited(self) -> None:
        self._new_deck()
        self.pg.keyboard.press("Control+m"); time.sleep(0.3)      # the new deck has the keyboard right away
        self.assertEqual(self.view("v.deck.slides.length"), 2)
        items = ["Черный чай проходит полную ферментацию листьев", "Зеленый чай не подвергается окислению при производстве",
                 "Белый чай изготавливается из молодых почек и листьев", "Улун имеет частичную степень ферментации от 10 до 80%",
                 "Пуэр проходит процесс микробиологического старения", "Каждый вид требует индивидуального подхода к завариванию",
                 "Седьмой пункт, чтобы текст точно не поместился в рамку", "Восьмой пункт для переполнения рамки текста"]
        self.pg.evaluate("(items) => { const v = MN.tabs.find(t => t.kind === 'slides').comp; v.change(() => { v.slide.elements[1].text.paras = items.map(t => ({ runs: [{ t }] })); }); }", items)
        time.sleep(0.3)
        r = self.pg.evaluate("""() => { const t = document.querySelectorAll('.sl-canvas .sl-slide > .sl-el')[1].querySelector('.sl-text');
            const ps = [...t.querySelectorAll('.sl-p')].map(p => p.getBoundingClientRect());
            const overlap = ps.slice(1).some((b, i) => b.top < ps[i].bottom - 1);
            const tb = t.getBoundingClientRect();
            return { fit: t.dataset.fit || null, overlap, bottom: ps.at(-1).bottom, boxBottom: tb.bottom }; }""")
        self.assertFalse(r["overlap"], r)
        self.assertEqual(r["fit"], "1")                 # layout body has no own autofit → the master's normAutofit applies
        self.assertLessEqual(r["bottom"], r["boxBottom"] + 2, r)
        # Home / End jump to the first / last slide (no selection)
        self.pg.keyboard.press("Escape"); self.pg.click(".sl-stage", position={"x": 5, "y": 5})
        self.pg.keyboard.press("Home"); time.sleep(0.1); self.assertEqual(self.view("v.cur"), 0)
        self.pg.keyboard.press("End"); time.sleep(0.1); self.assertEqual(self.view("v.cur"), 1)
        self.assertEqual(self.errors, [])

    def test_show_rapid_presses_during_transition_leave_one_slide(self) -> None:
        self._new_deck()
        self.js("""const v = MN.tabs.find(t => t.kind === 'slides').comp; for (let i = 0; i < 3; i++) v.addSlide();
          v.deck.slides.forEach((s, i) => { s.transition = { type: ['fade', 'push', 'wipe', 'circle'][i], dir: 'l', dur: 3000 }; v.touch(s); }); v.snapshot();
          localStorage.setItem('mn_sl_presenter', '0'); await v.startShow(0);""")
        self.pg.wait_for_selector(".sl-show")
        self.pg.keyboard.press("ArrowRight"); time.sleep(0.2)
        self.assertEqual(self.pg.evaluate("document.querySelectorAll('.sl-show-stage > *').length"), 2)   # old + new mid-transition
        self.pg.keyboard.press("ArrowRight"); time.sleep(0.2); self.pg.keyboard.press("ArrowRight"); time.sleep(0.2)
        self.assertLessEqual(self.pg.evaluate("document.querySelectorAll('.sl-show-stage > *').length"), 2)
        self.assertIn("4", self.pg.inner_text(".sl-hud-count"))
        time.sleep(3.3)
        self.assertEqual(self.pg.evaluate("document.querySelectorAll('.sl-show-stage > *').length"), 1)
        self.pg.keyboard.press("Escape"); time.sleep(0.2)
        self.assertIsNone(self.pg.query_selector(".sl-show"))
        self.assertEqual(self.errors, [])

    def _wait_js(self, expr: str, timeout: int = 20000) -> None:
        """Round 34: Page.wait_for_function eval()s its string, which MemuNote's CSP (script-src 'self') refuses;
        page.evaluate goes through CDP and is not subject to the page CSP."""
        end = time.time() + timeout / 1000
        while True:
            try:
                if self.pg.evaluate(f"({expr})()"):
                    return
            except Exception:
                pass
            if time.time() > end:
                raise AssertionError(f"timed out waiting for {expr}")
            time.sleep(0.1)

    def test_picture_copy_between_decks_and_pdf_export_with_image(self) -> None:
        """Copy/export read media from the model: fetch(blob:) is blocked by the page CSP (connect-src 'self')."""
        console = []
        self.pg.on("console", lambda m: m.type == "error" and console.append(m.text))
        self._new_deck()
        self.js("""const v = MN.tabs.find(t => t.kind === 'slides').comp; const c = document.createElement('canvas'); c.width = 64; c.height = 48;
          const x = c.getContext('2d'); x.fillStyle = '#e07a2f'; x.fillRect(0, 0, 64, 48);
          const b = await new Promise(r => c.toBlob(r, 'image/png')); await v.insertImageFile(new File([b], 'a.png', { type: 'image/png' }));""")
        time.sleep(0.3)
        self.view("v.el.querySelector('.sl-stage').focus()")   # Round 34: copy/paste are DOM events on the focused deck (a user clicks the slide first)
        self.pg.keyboard.press("Control+c"); time.sleep(0.5)
        pdf = self.dir / "img.pdf"
        self.pg.click("[data-c='rtab:show']"); self.pg.click("[data-c='exp:pdf']")
        self.pg.fill("#mn-prompt", str(pdf)); self.pg.keyboard.press("Enter")
        for _ in range(60):
            if pdf.exists() and pdf.stat().st_size > 1000:
                break
            time.sleep(0.25)
        self.assertTrue(pdf.exists() and pdf.read_bytes().startswith(b"%PDF"))
        self._new_deck()
        self.view("v.el.querySelector('.sl-stage').focus()")
        self.pg.keyboard.press("Control+v"); time.sleep(0.6)
        kinds = self.view("v.slide.elements.map(e => e.kind + ':' + !!(e.pic && v.deck.media.get(e.pic.media)))")
        self.assertIn("pic:true", kinds)
        self.assertEqual([c for c in console if "Content Security Policy" in c or "Failed to fetch" in c], [])
        self.assertEqual(self.errors, [])

    def test_statement_spec_keeps_its_text_and_tiny_charts_scale(self) -> None:
        r = self.js("""const P = await import('/js/slides/pptx.js'); const A = await import('/js/slides/slides-ai.js'); const R = await import('/js/slides/render.js');
          const deck = await P.newDeck({ lang: 'ru' });
          const txt = (s) => s.elements.filter(e => e.text).map(e => P.plainText(e.text)).join(' | ');
          const one = A.slideFromSpec(deck, { layout: 'statement', title: 'Главное правило', bullets: ['Свежая вода и точная температура'] });
          const many = A.slideFromSpec(deck, { layout: 'statement', title: 'Итог', bullets: ['Первое', 'Второе'] });
          const bare = A.slideFromSpec(deck, { layout: 'statement', title: 'Только мысль' });
          const svg = R.chartSvg({ type: 'col', cats: ['A', 'B'], series: [{ name: 'Продажи', vals: [1, 2] }] }, 120, 70, { accent: () => '#08c', color: () => '#08c' }, 1.5, '#000');
          return { one: txt(one), many: txt(many), bare: txt(bare), vb: svg.match(/viewBox="0 0 ([0-9.]+) /)[1], width: svg.match(/width="([0-9.]+)"/)[1] };""")
        self.assertIn("Свежая вода", r["one"]); self.assertIn("Второе", r["many"]); self.assertEqual(r["bare"], "Только мысль")
        self.assertEqual(float(r["width"]), 120.0); self.assertGreater(float(r["vb"]), 120 * 3)   # labels drawn at a readable virtual size

    def test_large_deck_opens_fast_and_thumbnails_are_lazy(self) -> None:
        r = self.js("""
          const P = await import('/js/slides/pptx.js');
          const deck = await P.newDeck({ lang: 'en' });
          for (let i = 0; i < 150; i++) { const s = P.makeSlide(deck, deck.layouts[1].path); s.elements[0].text = { paras: [{ runs: [{ t: 'Slide ' + (i + 2) }] }] }; s.elements[1].text = { paras: [{ runs: [{ t: 'Point ' + i }] }, { runs: [{ t: 'More text for the slide' }] }] }; deck.slides.push(s); }
          const t0 = performance.now(); const bytes = await P.savePptx(deck); const tSave = performance.now() - t0;
          const t1 = performance.now(); await P.openPptx(bytes); const tOpen = performance.now() - t1;
          let bin = ''; for (let i = 0; i < bytes.length; i += 8192) bin += String.fromCharCode.apply(null, bytes.subarray(i, i + 8192));
          return { b64: btoa(bin), tSave, tOpen };
        """)
        import base64
        big = self.dir / "big.pptx"
        big.write_bytes(base64.b64decode(r["b64"]))
        self.assertLess(r["tSave"], 15000)
        self.assertLess(r["tOpen"], 10000)
        t0 = time.time()
        self.open(big)
        self.assertEqual(self.view("v.deck.slides.length"), 151)
        self.assertLess(time.time() - t0, 15)
        painted = self.pg.evaluate("() => document.querySelectorAll('.sl-tbox .sl-slide').length")
        self.assertLess(painted, 60)                    # only visible thumbnails are rendered
        self.pg.click(".sl-side"); self.pg.keyboard.press("End"); time.sleep(0.6)
        self.assertEqual(self.view("v.cur"), 150)
        self.assertIn("Slide 151", self.pg.inner_text(".sl-canvas"))
        self.assertEqual(self.errors, [])

    def test_textedit_helpers(self) -> None:
        r = self.js("""
          const TE = await import('/js/slides/textedit.js');
          const paras = [{ runs: [{ t: 'Hello ', b: true }, { t: 'world' }] }, { lvl: 1, runs: [{ t: 'Second' }] }];
          const covered = TE.splitRange(paras, { start: { p: 0, o: 3 }, end: { p: 1, o: 3 } });
          TE.setRunProps(covered, { i: true });
          const after = JSON.parse(JSON.stringify(paras));
          const caret = TE.insertText(paras, { start: { p: 0, o: 0 }, end: { p: 0, o: 0 }, collapsed: true }, 'A\\nB');
          const w = TE.wordRange([{ runs: [{ t: 'one two three' }] }], { p: 0, o: 5 });
          return { after, caret, n: paras.length, w };
        """)
        self.assertEqual([[x["t"], bool(x.get("b")), bool(x.get("i"))] for x in r["after"][0]["runs"]], [["Hel", True, False], ["lo ", True, True], ["world", False, True]])
        self.assertEqual([[x["t"], bool(x.get("i"))] for x in r["after"][1]["runs"]], [["Sec", True], ["ond", False]])
        self.assertEqual(r["caret"], {"p": 1, "o": 1})
        self.assertEqual(r["n"], 3)
        self.assertEqual((r["w"]["start"]["o"], r["w"]["end"]["o"]), (4, 7))


# ============================================================================ live AI (real model)
@unittest.skipUnless(os.environ.get("XKIRO_BASE") and os.environ.get("XKIRO_KEY"), "live keys are not set")
class LiveGenerateTest(_ServerCase):
    """qwen3.8-max answers the deck prompt (through /api/slides_ai and an OpenAI-compatible relay) with parseable JSON."""

    def test_generate_deck_live(self) -> None:
        base = os.environ["XKIRO_BASE"].rstrip("/")
        keys = [os.environ.get(k) for k in ("XKIRO_KEY", "XKIRO_KEY2", "XKIRO_KEY3") if os.environ.get(k)]

        class Relay(http.server.BaseHTTPRequestHandler):
            def log_message(self, *a):  # noqa: D401
                pass

            def do_GET(self):  # noqa: N802
                self.send_response(200); self.end_headers(); self.wfile.write(b'{"models": []}')

            def do_POST(self):  # noqa: N802
                if not self.path.startswith("/v1/"):
                    self.send_response(404); self.end_headers(); return
                body = json.loads(self.rfile.read(int(self.headers.get("Content-Length") or 0)))
                body["model"] = "qwen/qwen3.8-max:free"
                last = None
                for key in keys:
                    req = urllib.request.Request(base + "/chat/completions", data=json.dumps(body).encode(), method="POST",
                                                 headers={"Content-Type": "application/json", "Authorization": "Bearer " + key, "User-Agent": "MemuNote-tests/3.5"})
                    try:
                        resp = urllib.request.urlopen(req, timeout=240)
                        self.send_response(200); self.send_header("Content-Type", "text/event-stream"); self.end_headers()
                        while True:
                            chunk = resp.read(1024)
                            if not chunk:
                                break
                            self.wfile.write(chunk)
                        return
                    except Exception as exc:  # noqa: BLE001 - 429 → next key
                        last = exc
                self.send_response(502); self.end_headers(); self.wfile.write(str(last).encode()[:300])

        relay = http.server.ThreadingHTTPServer(("127.0.0.1", 0), Relay)
        threading.Thread(target=relay.serve_forever, daemon=True).start()
        try:
            with mock.patch.dict(os.environ, {"MEMUAI_PROXY_URL": f"http://127.0.0.1:{relay.server_address[1]}"}):
                _, raw = self.call("POST", "/api/slides_ai", {"task": "generate", "model": "qwen/qwen3.8-max:free",
                                                              "input": {"topic": "Как заваривать чай: виды, температура, время", "count": 5, "lang": "ru"}}, raw=True)
        finally:
            relay.shutdown()
        events = [json.loads(line[5:]) for line in raw.decode("utf-8").split("\n") if line.startswith("data:")]
        err = [e for e in events if e["t"] == "error"]
        if err and re.search(r"429|rate|quota", err[0].get("message", ""), re.I):
            self.skipTest("model quota: " + err[0]["message"][:200])
        self.assertFalse(err, err)
        text = "".join(e["text"] for e in events if e["t"] == "delta")
        m = re.search(r"\{[\s\S]*\}", text)
        self.assertTrue(m, text[:500])
        deck = json.loads(m.group(0))
        self.assertGreaterEqual(len(deck["slides"]), 4)
        self.assertEqual(deck["slides"][0]["layout"], "title")
        self.assertTrue(re.search(r"[А-Яа-я]", json.dumps(deck, ensure_ascii=False)))


if __name__ == "__main__":
    unittest.main()
